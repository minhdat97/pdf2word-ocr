"""End-to-end test against a running server (default http://127.0.0.1:8000).

Runs the full upload -> analyse -> convert -> download flow for every sample,
checks the text against the born-digital source, checks the formatting the
native pipeline must keep (exam-like.pdf), then probes error paths.
"""

import difflib
import io
import json
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.request
import uuid
import zipfile
from pathlib import Path

import docx
import pymupdf

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000"
ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
sys.path.insert(0, str(ROOT))
from app.vnfont import decode_span  # noqa: E402


def call(method, path, body=None, raw=None, headers=None):
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    hdrs = {"content-type": "application/json"} if body is not None else {}
    req = urllib.request.Request(BASE + path, data=data, method=method, headers={**hdrs, **(headers or {})})
    try:
        with urllib.request.urlopen(req) as res:
            return res.status, res.read()
    except urllib.error.HTTPError as err:
        return err.code, err.read()


def upload(data: bytes, name="file.pdf") -> str:
    status, body = call("POST", "/v1/upload-url", {"filename": name})
    assert status == 200, body
    up = json.loads(body)
    status, body = call("PUT", up["url"], raw=data, headers={"content-type": "application/pdf"})
    assert status == 200, body
    return up["token"]


def run_task(token, tool, options=None, timeout=300):
    task_id = uuid.uuid4().hex
    status, body = call(
        "POST",
        "/v1/tasks",
        {"tool": tool, "task_id": task_id, "input_tokens": [token], "options": options or {}},
    )
    assert status == 201, body
    deadline = time.time() + timeout
    while time.time() < deadline:
        data = json.loads(call("GET", f"/v1/tasks/{task_id}")[1])["data"]
        if data["status"] in ("done", "failed"):
            return data["result"]
        time.sleep(0.3)
    raise TimeoutError(task_id)


def docx_text(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for section in d.sections:
        parts.extend(p.text for p in section.header.paragraphs + section.footer.paragraphs)
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return " ".join(parts)


def pdf_text(path: Path) -> str:
    """The PDF's text layer, with legacy VNI/TCVN3 spans decoded."""
    parts = []
    for page in pymupdf.open(path):
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                if abs(line["dir"][1]) > 0.02:
                    continue  # rotated stamps are kept as background pictures
                parts.extend(decode_span(s["text"], s["font"])[0] for s in line["spans"])
                parts.append(" ")
    return "".join(parts).replace("\xad", "-")  # the converter prints soft hyphens


def formatting_checks(data: bytes) -> list[tuple[str, bool]]:
    z = zipfile.ZipFile(io.BytesIO(data))
    body = z.read("word/document.xml").decode()
    parts = {n: z.read(n).decode() for n in z.namelist() if n.startswith(("word/header", "word/footer"))}
    headers = " ".join(v for k, v in parts.items() if "header" in k)
    footers = " ".join(v for k, v in parts.items() if "footer" in k)
    return [
        ("VNI title decoded to Unicode", "ĐỀ KIỂM TRA SỐ 1" in body),
        ("superscripts kept", body.count('w:val="superscript"') >= 5),
        ("subscript kept", 'w:val="subscript"' in body),
        ("colours kept", 'w:color w:val="0000FF"' in body and 'w:color w:val="FF0000"' in body),
        ("answer options on tab stops", body.count("<w:tab/>") >= 3 and "<w:tabs>" in body),
        ("figure floated on the page", "<wp:anchor" in body),
        ("rotated stamp kept behind the text", 'behindDoc="1"' in body),
        ("letterhead as side-by-side cells", "<w:tbl>" in body),
        # Header/footer hold "Trang x/y" only; what repeated there is gone.
        (
            "repeated header/footer text removed",
            not any(t in body + headers + footers for t in ("Trường THPT Mẫu", "Đề thi thử", "dethi.example.vn")),
        ),
        (
            "footer is 'Trang PAGE/NUMPAGES'",
            "Trang " in footers and 'w:instr="PAGE ' in footers and 'w:instr="NUMPAGES ' in footers,
        ),
    ]


def _watermark_pixels(z: zipfile.ZipFile) -> int:
    """Pale blue pixels (the sample's watermark colour over white) in the DOCX images."""
    import cv2
    import numpy as np

    count = 0
    for name in z.namelist():
        if name.startswith("word/media/"):
            img = cv2.imdecode(np.frombuffer(z.read(name), np.uint8), cv2.IMREAD_COLOR)
            if img is not None:
                blue, red = img[..., 0].astype(int), img[..., 2].astype(int)
                count += int(((blue - red >= 6) & (img.min(axis=2) > 200)).sum())
    return count


def norm(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", text).split())


def similarity(a: str, b: str) -> float:
    # Order-insensitive: compare sorted words, since tables come after paragraphs in docx_text.
    return difflib.SequenceMatcher(
        None, sorted(norm(a).split()), sorted(norm(b).split()), autojunk=False
    ).ratio()


def main():
    truth_vi = pdf_text(SAMPLES / "vi-text.pdf")
    cases = [
        ("vi-scan.pdf", "scan", {"ocr_type": "ocr", "languages": ["vie"]}, truth_vi),
        ("mixed.pdf", "mixed", {"ocr_type": "ocr", "languages": ["vie"]}, truth_vi),
        ("vi-text.pdf", "text", {"ocr_type": "", "languages": []}, truth_vi),
        ("vi-text.pdf", "text", {"ocr_type": "ocr-full", "languages": ["vie"]}, truth_vi),
        ("en-scan.pdf", "scan", {"ocr_type": "ocr", "languages": ["eng"]}, None),
        (
            "exam-like.pdf",
            "text",
            {"ocr_type": "ocr", "languages": ["vie"]},
            # header/footer lines are dropped by design (only "Trang x/y" is kept)
            re.sub(r"Trường\s+THPT\s+Mẫu|Nguồn:\s+kho\s+đề\s+mẫu|Đề\s+thi\s+thử|https://dethi\.example\.vn|Trang\s+\d/2", " ",
                   unicodedata.normalize("NFC", pdf_text(SAMPLES / "exam-like.pdf"))),
        ),
        # Scans: straightened before OCR, figures kept as pictures.
        ("vi-scan-skewed.pdf", "scan", {"ocr_type": "ocr", "languages": ["vie"]}, truth_vi),
        ("vi-scan-turned.pdf", "scan", {"ocr_type": "ocr", "languages": ["vie"]}, truth_vi),
        ("exam-like-scan.pdf", "scan", {"ocr_type": "ocr", "languages": ["vie"]}, None),
    ]
    failures = 0
    exam_docx = scan_docx = None
    for name, want_kind, options, truth in cases:
        t0 = time.time()
        token = upload((SAMPLES / name).read_bytes(), name)
        analysis = run_task(token, "analyse-document")
        kind = analysis["options"]["kind"]
        result = run_task(token, "pdf-to-word", options)
        assert result["success"], result
        status, data = call("GET", f"/v1/files/{result['output_files'][0]}?name=out.docx")
        assert status == 200 and data[:2] == b"PK", status
        text = docx_text(data)
        if name == "exam-like.pdf":
            exam_docx = data
        if name == "exam-like-scan.pdf":
            scan_docx = data
        score = similarity(truth, text) if truth else None
        ok = kind == want_kind and (score is None or score > 0.9)
        if name == "exam-like-scan.pdf":
            ok = ok and result["options"]["pictures"] >= 1  # the lever figure survives OCR
        failures += not ok
        info = result["options"]
        print(
            f"{'OK ' if ok else 'BAD'} {name:13} {info['mode']:8} kind={kind:5} "
            f"pages native/ocr/image={info['pages_native']}/{info['pages_ocr']}/{info['pages_image']} "
            f"conf={info.get('avg_confidence')} similarity={score if score is None else round(score, 3)} "
            f"{time.time() - t0:.1f}s"
        )

    for label, passed in formatting_checks(exam_docx):
        failures += not passed
        print(f"{'OK ' if passed else 'BAD'} exam-like: {label}")

    # Scanned page numbers ("Trang 1/2" read by OCR) become Word fields, not static text.
    z = zipfile.ZipFile(io.BytesIO(scan_docx))
    footers = " ".join(z.read(n).decode() for n in z.namelist() if n.startswith("word/footer"))
    static = re.findall(r"<w:t[^>]*>\s*Trang\s*\d", z.read("word/document.xml").decode())
    passed = 'w:instr="PAGE ' in footers and not static
    failures += not passed
    print(f"{'OK ' if passed else 'BAD'} exam-like-scan: page numbers are Word fields in the footer")

    # strip_watermarks: stamps, background watermarks and links go; the page
    # numbers stay (the only header/footer content kept in any mode).
    token = upload((SAMPLES / "exam-like.pdf").read_bytes())
    result = run_task(token, "pdf-to-word", {"ocr_type": "", "strip_watermarks": True})
    data = call("GET", f"/v1/files/{result['output_files'][0]}")[1]
    z = zipfile.ZipFile(io.BytesIO(data))
    body = z.read("word/document.xml").decode()
    rest = " ".join(z.read(n).decode() for n in z.namelist() if n.startswith(("word/header", "word/footer")))
    for label, passed in [
        ("background watermark removed", 'behindDoc="1"' not in body),
        ("figure kept", "<wp:anchor" in body),
        ("footer link removed", "dethi.example.vn" not in body + rest),
        ("translucent stamp removed", "kho đề mẫu" not in body + rest),
        ("page numbers kept", 'w:instr="PAGE ' in rest),
        # ~20k watermark pixels show through the figure unless cleared; only
        # anti-aliased edges of the figure may keep a tint.
        ("figure freed of the watermark showing through it", _watermark_pixels(z) < 1000),
        (
            "removals reported",
            result["options"]["removed_watermarks"] >= 1 and result["options"]["removed_lines"] >= 2,
        ),
    ]:
        failures += not passed
        print(f"{'OK ' if passed else 'BAD'} strip_watermarks: {label}")

    # Error paths.
    up = json.loads(call("POST", "/v1/upload-url", {"filename": "x.pdf"})[1])
    checks = [
        ("tampered signature -> 403", call("PUT", up["url"][:-4] + "beef", raw=b"%PDF-1.4")[0] == 403),
        ("non-PDF body -> 400", call("PUT", up["url"], raw=b"hello")[0] == 400),
        (
            "bad task_id -> 400",
            call("POST", "/v1/tasks", {"tool": "pdf-to-word", "task_id": "x", "input_tokens": []})[0] == 400,
        ),
        (
            "unknown token -> 400",
            call(
                "POST",
                "/v1/tasks",
                {
                    "tool": "pdf-to-word",
                    "task_id": uuid.uuid4().hex,
                    "input_tokens": [uuid.uuid4().hex + ".pdf"],
                },
            )[0]
            == 400,
        ),
        (
            "unknown language -> 400",
            call(
                "POST",
                "/v1/tasks",
                {
                    "tool": "pdf-to-word",
                    "task_id": uuid.uuid4().hex,
                    "input_tokens": [upload((SAMPLES / "en-scan.pdf").read_bytes())],
                    "options": {"ocr_type": "ocr", "languages": ["klingon"]},
                },
            )[0]
            == 400,
        ),
    ]
    enc = pymupdf.open(SAMPLES / "vi-text.pdf")
    buf = io.BytesIO()
    enc.save(buf, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="secret")
    token = upload(buf.getvalue())
    checks.append(
        (
            "password PDF -> analysis flags it",
            run_task(token, "analyse-document")["options"]["password_protected"] is True,
        )
    )
    checks.append(
        (
            "password PDF -> convert error_fe=password",
            run_task(token, "pdf-to-word", {"ocr_type": "ocr"})["error_fe"] == "password",
        )
    )
    # The page closed while its conversion runs (POST .../cancel, sent by the page): it stops at its next page.
    token = upload((SAMPLES / "vi-scan.pdf").read_bytes())
    task_id = uuid.uuid4().hex
    call("POST", "/v1/tasks", {"tool": "pdf-to-word", "task_id": task_id, "input_tokens": [token],
                               "options": {"ocr_type": "ocr"}})
    cancelled = call("POST", f"/v1/tasks/{task_id}/cancel")[0] == 200
    deadline = time.time() + 120
    while time.time() < deadline:
        data = json.loads(call("GET", f"/v1/tasks/{task_id}")[1])["data"]
        if data["status"] in ("done", "failed"):
            break
        time.sleep(0.3)
    checks.append(("cancel -> task stopped, error_fe=cancelled",
                   cancelled and data["status"] == "failed" and data["result"]["error_fe"] == "cancelled"))
    checks.append(("cancel unknown task -> 404", call("POST", f"/v1/tasks/{uuid.uuid4().hex}/cancel")[0] == 404))
    for label, passed in checks:
        failures += not passed
        print(f"{'OK ' if passed else 'BAD'} {label}")

    print("FAILED" if failures else "ALL PASSED", f"({failures} failures)")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
