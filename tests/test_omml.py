"""Tests for app.omml (LaTeX -> Word equation). Run: python -m unittest discover -s tests"""

import re
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from lxml import etree  # noqa: E402

from app.omml import M_NS, W_NS, latex_to_omml  # noqa: E402

M = f"{{{M_NS}}}"
W = f"{{{W_NS}}}"

# Formulas of the 3-page physics exam used to tune the OCR path (its answer options).
GOLD = [
    'p',
    '\\rho',
    'v_{cqp}=\\sqrt{\\frac{3p}{\\rho}}',
    'v_{cqp}=\\sqrt{\\frac{p}{\\rho}}',
    'v_{cqp}=\\sqrt{\\frac{p}{3\\rho}}',
    'v_{cqp}=\\sqrt{\\frac{3p}{2\\rho}}',
    '2\\sqrt{3}',
    'k',
    'T',
    'm_{0}',
    'v_{cqp}=\\sqrt{\\frac{2kT}{3m_{0}}}',
    'v_{cqp}=\\sqrt{\\frac{kT}{3m_{0}}}',
    'v_{cqp}=\\sqrt{\\frac{3kT}{m_{0}}}',
    'v_{cqp}=\\sqrt{\\frac{3kT}{2m_{0}}}',
    '400\\sqrt{2}',
    'n_{1}',
    'T_{1}',
    'n_{2}',
    'T_{2}',
    '\\frac{n_{1}T_{1}+n_{2}T_{2}}{n_{1}+n_{2}}',
    '\\frac{n_{1}T_{2}-n_{2}T_{1}}{n_{1}-n_{2}}',
    '\\frac{n_{1}T_{1}-n_{2}T_{2}}{n_{1}-n_{2}}',
    '\\frac{n_{1}T_{2}+n_{2}T_{1}}{n_{1}+n_{2}}',
]

# What the formula model read from the option crops of that exam, verbatim (spaced tokens,
# misreads and all): every one must convert.
MODEL = [
    '\\mathrm { n h i \\hat { e } t ~ d \\hat { o } ~ c u a ~ k h i . }',
    '\\mathrm { t h } \\hat { \\hat { e } } \\mathrm { ~ t i c h ~ c u a ~ k h i . }',
    '\\acute { \\mathrm { a p } } \\ \\mathrm { s u \\acute { a } } \\mathrm { t } \\ \\mathrm { c i a \\ k h i } .',
    '\\mathrm { m \\hat { a } t \\ d \\hat { o } \\ p h \\hat { a } n \\ t i r . }',
    '6 , 4 2 . 1 0 ^ { - 2 1 } \\: \\mathrm { J } .',
    '\\mathrm { 7 , 6 6 . 1 0 ^ { \\theta ^ { 2 2 } } \\, J . }',
    '\\mathrm { 1 , 2 3 . 1 0 ^ { - 2 1 } ~ J . }',
    '8 , 2 1 . 1 0 ^ { - 2 2 } \\, \\mathrm { J } .',
    '5 , 3 2 . 1 0 ^ { - 2 1 } \\: \\mathrm { J } .',
    '6 , \\! 2 3 . 1 0 ^ { - \\! 2 1 } \\: \\mathrm { J } \\! -',
    "\\mathrm { 7 , 4 3 . 1 0 ^ { ' 2 1 } \\: J . }",
    '4 , 5 7 . 1 0 ^ { - 2 1 } \\: \\mathrm { J } .',
    '6 \\mathrm { \\ k P a }',
    '3 \\, \\, \\mathrm { k P a . }',
    '5 \\mathrm { ~ k P a . }',
    '1 5 \\, \\, \\mathrm { k P a . }',
    'v _ { { } _ { c q p } } ^ { { } } = \\sqrt { \\frac { 3 p } { \\rho } } ~ .',
    'v _ { { } _ { c q p } } = \\sqrt { \\frac { p } { \\rho } } \\ .',
    'v _ { { } _ { c q p } } ^ { { } } = \\sqrt { \\frac { p } { 3 \\rho } } ~ .',
    'v _ { { } _ { c q p } } ^ { { } } = \\sqrt { \\frac { 3 p } { 2 \\rho } } \\ .',
    '\\mathrm { t a n g ~ 9 ~ l a n } .',
    '\\mathrm { t a n g ~ 3 ~ \\hat { l a n } . }',
    '\\mathrm { g i a m } \\; 3 \\; \\mathrm { l a n } .',
    '\\mathrm { \\ g i a m \\ 9 \\ l a n } .',
    '\\mathrm { t / a n g ~ \\, 2 \\sqrt 3 ~ \\, ~ \\dot { l a n } . }',
    '\\mathrm { t a n g ~ 1 2 ~ l a n . }',
    '\\mathrm { \\ g i a m \\ } 2 \\sqrt { 3 } \\ \\mathrm { l \\hat { a } n } .',
    '\\mathrm { g i a m ~ 1 2 ~ l a n } .',
    '5 9 4 ~ \\mathrm { m / s . }',
    '3 4 3 ~ \\mathrm { m / s . }',
    '\\mathrm { 4 2 0 ~ m / s . }',
    '2 8 0 ~ \\mathrm { m / s . }',
    '=',
    '\\mathbf { \\psi } _ { \\mathbf { } _ { c q p } } = \\sqrt { \\frac { 3 k T } { m _ { 0 } } } \\ .',
    '\\boldsymbol { v } _ { { } _ { c q p } } = \\sqrt { \\frac { 3 k T } { 2 m _ { 0 } } } \\ .',
    '0 , 5 . 1 0 ^ { 3 } ~ \\mathrm { m / s . }',
    '\\mathrm { 1 , 4 . 1 0 ^ { 3 } ~ m / s . }',
    '\\mathrm { 1 , 1 . 1 0 ^ { 3 } ~ m / s . }',
    '\\mathrm { 1 , } 9 . 1 0 ^ { 3 } ~ \\mathrm { m / s . }',
    '2 3 6 ~ \\mathrm { m / s . }',
    '\\mathrm { 1 6 7 ~ m / s . }',
    '3 5 4 ~ \\mathrm { m / s . }',
    '5 0 0 ~ \\mathrm { m / s . }',
    '1 , 2 2 .',
    '1 , 4 8 .',
    '2 , 8 3 .',
    '8 , 3 \\; \\mathrm { ^ \\circ C . }',
    '1 4 , 1 \\, \\, ^ { \\circ } \\mathrm { C } .',
    '1 9 , 2 \\, \\, ^ { \\circ } \\mathrm { C } .',
    '1 9 , 2 ~ \\mathrm { m / s . }',
    '\\mathrm { 1 9 2 ~ m / s . }',
    '\\mathrm { 1 , } 9 2 . 1 0 ^ { 4 } \\mathrm { ~ m / s . }',
    '\\mathrm { 1 , } 9 2 . 1 0 ^ { 5 } ~ \\mathrm { m / s . }',
    '8 ~ \\mathrm { k m / s . }',
    '2 ~ \\mathrm { k m / s . }',
    '1 \\, \\mathrm { \\ k m / s . }',
    '0 , 5 ~ \\mathrm { k m / s . }',
    '1 8 9 ~ \\mathrm { ^ \\circ C . }',
    '1 3 2 3 \\, \\, \\circ \\mathrm { C } .',
    '5 2 1 \\, \\, \\mathrm { ^ \\circ C . }',
    '1 8 2 7 \\, \\, { } ^ { \\circ } \\mathrm { C } .',
    '4 _ { \\mathrm { V O } } = 9 _ { \\mathrm { V H } } .',
    '9 _ { \\mathrm { V O } } = 1 3 4 _ { \\mathrm { V H } } .',
    '\\mathrm { \\Phi } _ { \\mathrm { V O } } = 3 \\mathrm { v } _ { \\mathrm { H } } .',
    '7 2 \\mathrm { v _ { O } } = 6 7 \\mathrm { v _ { H } } .',
    '8 0 0 ~ \\mathrm { m / s . }',
    '4 0 0 ~ \\mathrm { m / s . }',
    '4 0 0 { \\sqrt 2 } ~ \\mathrm { m / s . }',
    '2 0 0 ~ \\mathrm { m / s . }',
    '3 2 ~ \\mathrm { ^ \\circ C . }',
    '2 7 ~ \\mathrm { ^ \\circ C . }',
    '3 0 , 5 \\; \\mathrm { ^ \\circ C . }',
]

# Typical high-school physics and maths formulas.
WRITTEN = [
    r"pV=nRT",
    r"\frac{pV}{T}=\text{hằng số}",
    r"p=\frac{1}{3}\mu n m_{0}\overline{v^{2}}",
    r"\overline{W_{đ}}=\frac{3}{2}kT",
    r"k=\frac{R}{N_{A}}=1{,}38.10^{-23}\,\mathrm{J/K}",
    r"T=t+273",
    r"\Delta U=A+Q",
    r"Q=mc\Delta t",
    r"\frac{p_{1}V_{1}}{T_{1}}=\frac{p_{2}V_{2}}{T_{2}}",
    r"v_{cqp}=\sqrt{\frac{3RT}{\mu}}",
    r"\rho=\frac{m}{V}",
    r"I=\frac{U}{R}",
    r"I=\frac{\mathcal{E}}{R_{N}+r}",
    r"R=\rho\frac{l}{S}",
    r"P=UI=I^{2}R=\frac{U^{2}}{R}",
    r"F=k\frac{|q_{1}q_{2}|}{\varepsilon r^{2}}",
    r"\vec{F}=q\vec{E}",
    r"\Phi=BS\cos\alpha",
    r"e_{c}=-\frac{\Delta\Phi}{\Delta t}",
    r"Z=\sqrt{R^{2}+\left(Z_{L}-Z_{C}\right)^{2}}",
    r"\omega=\frac{1}{\sqrt{LC}}",
    r"u=U_{0}\cos(\omega t+\varphi)",
    r"\tan\varphi=\frac{Z_{L}-Z_{C}}{R}",
    r"B=4\pi.10^{-7}\frac{NI}{l}",
    r"v=v_{0}+at",
    r"s=v_{0}t+\frac{1}{2}at^{2}",
    r"v^{2}-v_{0}^{2}=2as",
    r"\vec{a}=\frac{\Delta\vec{v}}{\Delta t}",
    r"F_{hd}=G\frac{m_{1}m_{2}}{r^{2}}",
    r"W_{đ}=\frac{1}{2}mv^{2}",
    r"T=2\pi\sqrt{\frac{l}{g}}",
    r"x=A\cos\left(\omega t+\varphi\right)",
    r"\lambda=\frac{v}{f}=vT",
    r"\Delta t=\frac{\Delta t_{0}}{\sqrt{1-\frac{v^{2}}{c^{2}}}}",
    r"t=27^{\circ}\mathrm{C}",
    r"\eta=\frac{A}{Q_{1}}.100\%",
    r"{}_{92}^{235}\mathrm{U}+{}_{0}^{1}n\to{}_{56}^{141}\mathrm{Ba}+{}_{36}^{92}\mathrm{Kr}+3{}_{0}^{1}n",
    r"N=N_{0}e^{-\lambda t}",
    r"T=\frac{\ln 2}{\lambda}",
    r"\log_{a}b=\frac{\ln b}{\ln a}",
    r"\log_{2}(x-1)\le 3",
    r"\lim_{x\to 0}\frac{\sin x}{x}=1",
    r"\lim_{x\to+\infty}\left(1+\frac{1}{x}\right)^{x}=e",
    r"\sum_{k=1}^{n}k=\frac{n(n+1)}{2}",
    r"\int_{0}^{1}x^{2}\,dx=\frac{1}{3}",
    r"\int\limits_{a}^{b}f(x)\,\mathrm{d}x=F(b)-F(a)",
    r"\left.F(x)\right|_{a}^{b}",
    r"x=\frac{-b\pm\sqrt{b^{2}-4ac}}{2a}",
    r"\sqrt[3]{x^{2}}=x^{\frac{2}{3}}",
    r"\sin^{2}x+\cos^{2}x=1",
    r"y'=3x^{2}-6x",
    r"f''(x)>0",
    r"\overrightarrow{AB}+\overrightarrow{BC}=\overrightarrow{AC}",
    r"\left|\vec{a}\right|=\sqrt{x^{2}+y^{2}}",
    r"\vec{a}\cdot\vec{b}=|\vec{a}|.|\vec{b}|\cos(\vec{a},\vec{b})",
    r"\begin{cases}x+y=3\\x-y=1\end{cases}",
    r"f(x)=\begin{cases}x^{2} & \text{khi } x\ge 0\\-x & \text{khi } x<0\end{cases}",
    r"\left\{\begin{array}{l}x=1+2t\\y=3-t\end{array}\right.",
    r"\left[\begin{array}{l}x=1\\x=2\end{array}\right.",
    r"x\in(0;+\infty)",
    r"x\in\left[1;3\right)",
    r"D=\mathbb{R}\setminus\{1\}",
    r"C_{n}^{k}=\frac{n!}{k!(n-k)!}",
    r"\binom{n}{k}=\dbinom{n}{n-k}",
    r"A=\begin{pmatrix}1&2\\3&4\end{pmatrix}",
    r"\det\begin{vmatrix}a&b\\c&d\end{vmatrix}=ad-bc",
    r"\widehat{ABC}=60^{\circ}",
    r"\Delta=b^{2}-4ac<0",
    r"e^{i\pi}+1=0",
    r"a_{n}=a_{1}+(n-1)d",
    r"S_{n}=\frac{n(a_{1}+a_{n})}{2}",
    r"\frac{\mathrm{d}y}{\mathrm{d}x}",
    r"\sqrt{x}\geq 0\Rightarrow x\neq -1",
    r"y=\frac{2x+1}{x-1}",
    r"\lg 100=2",
    r"\exp(x)\approx 1+x",
    r"\overline{abc}=100a+10b+c",
    r"\operatorname{tg}\alpha=\frac{\sin\alpha}{\cos\alpha}",
    r"\prod_{i=1}^{n}a_{i}",
    r"\oint\vec{B}\cdot d\vec{l}=\mu_{0}I",
    r"\begin{aligned}x&=1\\y&=2\end{aligned}",
    r"\mathbf{F}=m\mathbf{a}",
    r"\hat{x}+\dot{x}+\ddot{x}+\tilde{x}+\bar{x}",
    r"\left\langle x\right\rangle",
    r"\left(\frac{a}{b}\middle|c\right)",
    r"a\times b\div c\mp d",
    r"x\to\infty\Leftrightarrow\frac{1}{x}\to 0",
    r"\{x\in\mathbb{R}\mid x>0\}",
    r"\left\lfloor x\right\rfloor+\lceil y\rceil",
    r"\Delta ABC\sim\Delta A'B'C'",
    r"\overset{t^{\circ}}{\rightarrow}",
    r"\begin{pmatrix}a&\begin{vmatrix}1&2\\3&4\end{vmatrix}\end{pmatrix}",
]

# Exact structure for representative formulas (see skeleton(): Pr elements dropped except the
# properties in SHOWN; r@p = upright run, r@b = bold, r~ = normal text).
SKELETONS = {
    r"v_{cqp}=\sqrt{\frac{3p}{\rho}}":
        "sSub[e[r:v] sub[r:cqp]] r:= rad[deg[] e[f[num[r:3p] den[r:ρ]]]]",
    r"\frac{n_{1}T_{2}-n_{2}T_{1}}{n_{1}-n_{2}}":
        "f[num[sSub[e[r:n] sub[r:1]] sSub[e[r:T] sub[r:2]] r:− sSub[e[r:n] sub[r:2]] sSub[e[r:T] sub[r:1]]] "
        "den[sSub[e[r:n] sub[r:1]] r:− sSub[e[r:n] sub[r:2]]]]",
    r"400\sqrt{2}": "r:400 rad[deg[] e[r:2]]",
    # model output: spaced tokens, \: spacing, decimal comma, upright unit
    "6 , 4 2 . 1 0 ^ { - 2 1 } \\: \\mathrm { J } .": "r:6,42. sSup[e[r:10] sup[r:−21]] r:\u00a0 r@p:J r:.",
    # empty bases and empty scripts the model writes around subscripts
    "v _ { { } _ { c q p } } ^ { { } } = \\sqrt { \\frac { 3 p } { \\rho } } ~ .":
        "sSub[e[r:v] sub[r:cqp]] r:= rad[deg[] e[f[num[r:3p] den[r:ρ]]]] r:.",
    # degree written three ways
    "1 8 2 7 \\, \\, { } ^ { \\circ } \\mathrm { C } .": "r:1827\u00a0° r@p:C r:.",
    "1 3 2 3 \\, \\, \\circ \\mathrm { C } .": "r:1323\u00a0° r@p:C r:.",
    "8 , 3 \\; \\mathrm { ^ \\circ C . }": "r:8,3\u00a0 r@p:°C.",
    "4 0 0 { \\sqrt 2 } ~ \\mathrm { m / s . }": "r:400 rad[deg[] e[r:2]] r:\u00a0 r@p:m/s.",
    # one space at most, none before punctuation, none at the ends (also inside groups)
    "1 \\, \\mathrm { \\ k m / s . }": "r:1\u00a0 r@p:km/s.",
    "\\mathrm { \\ g i a m \\ 9 \\ l a n } .": "r@p:giam\u00a09\u00a0lan r:.",
    "v _ { { } _ { c q p } } = \\sqrt { \\frac { p } { \\rho } } \\ .":
        "sSub[e[r:v] sub[r:cqp]] r:= rad[deg[] e[f[num[r:p] den[r:ρ]]]] r:.",
    r"t=27^{\circ}\mathrm{C}": "r:t=27° r@p:C",
    r"\frac{pV}{T}=\text{hằng số}": "f[num[r:pV] den[r:T]] r:= r~:hằng số",
    r"p=\frac{1}{3}\mu n m_{0}\overline{v^{2}}":
        "r:p= f[num[r:1] den[r:3]] r:μn sSub[e[r:m] sub[r:0]] bar{pos=top}[e[sSup[e[r:v] sup[r:2]]]]",
    r"\mathbf{F}=m\mathbf{a}": "r@b:F r:=m r@b:a",
    r"\vec{a}=\frac{\Delta\vec{v}}{\Delta t}":
        "acc{chr=\u20d7}[e[r:a]] r:= f[num[r:Δ acc{chr=\u20d7}[e[r:v]]] den[r:Δt]]",
    r"\hat{x}+\dot{x}+\ddot{x}+\tilde{x}+\bar{x}":
        "acc{chr=\u0302}[e[r:x]] r:+ acc{chr=\u0307}[e[r:x]] r:+ acc{chr=\u0308}[e[r:x]] r:+ "
        "acc{chr=\u0303}[e[r:x]] r:+ acc{chr=\u0305}[e[r:x]]",
    r"\sqrt[3]{x^{2}}=x^{\frac{2}{3}}":
        "rad[deg[r:3] e[sSup[e[r:x] sup[r:2]]]] r:= sSup[e[r:x] sup[f[num[r:2] den[r:3]]]]",
    r"y'=3x^{2}-6x": "sSup[e[r:y] sup[r:′]] r:=3 sSup[e[r:x] sup[r:2]] r:−6x",
    r"\Delta ABC\sim\Delta A'B'C'":
        "r:ΔABC∼Δ sSup[e[r:A] sup[r:′]] sSup[e[r:B] sup[r:′]] sSup[e[r:C] sup[r:′]]",
    r"v^{2}-v_{0}^{2}=2as": "sSup[e[r:v] sup[r:2]] r:− sSubSup[e[r:v] sub[r:0] sup[r:2]] r:=2as",
    # functions take the following operand
    r"\sin^{2}x+\cos^{2}x=1":
        "func[fName[sSup[e[r@p:sin] sup[r:2]]] e[r:x]] r:+ func[fName[sSup[e[r@p:cos] sup[r:2]]] e[r:x]] r:=1",
    r"u=U_{0}\cos(\omega t+\varphi)": "r:u= sSub[e[r:U] sub[r:0]] func[fName[r@p:cos] e[d[e[r:ωt+φ]]]]",
    r"\log_{2}(x-1)\le 3": "func[fName[sSub[e[r@p:log] sub[r:2]]] e[d[e[r:x−1]]]] r:≤3",
    r"\lim_{x\to 0}\frac{\sin x}{x}=1":
        "func[fName[limLow[e[r@p:lim] lim[r:x→0]]] e[f[num[func[fName[r@p:sin] e[r:x]]] den[r:x]]]] r:=1",
    r"\lim_{x\to+\infty}\left(1+\frac{1}{x}\right)^{x}=e":
        "func[fName[limLow[e[r@p:lim] lim[r:x→+∞]]] e[sSup[e[d[e[r:1+ f[num[r:1] den[r:x]]]]] sup[r:x]]]] r:=e",
    r"\operatorname{tg}\alpha=\frac{\sin\alpha}{\cos\alpha}":
        "func[fName[r@p:tg] e[r:α]] r:= f[num[func[fName[r@p:sin] e[r:α]]] den[func[fName[r@p:cos] e[r:α]]]]",
    # n-ary operators
    r"\sum_{k=1}^{n}k=\frac{n(n+1)}{2}":
        "nary{chr=∑,limLoc=undOvr}[sub[r:k=1] sup[r:n] e[r:k]] r:= f[num[r:n d[e[r:n+1]]] den[r:2]]",
    r"\int_{0}^{1}x^{2}\,dx=\frac{1}{3}":
        "nary{chr=∫,limLoc=subSup}[sub[r:0] sup[r:1] e[sSup[e[r:x] sup[r:2]] r:\u00a0dx]] r:= f[num[r:1] den[r:3]]",
    r"\int\limits_{a}^{b}f(x)\,\mathrm{d}x=F(b)-F(a)":
        "nary{chr=∫,limLoc=undOvr}[sub[r:a] sup[r:b] e[r:f d[e[r:x]] r:\u00a0 r@p:d r:x]] "
        "r:=F d[e[r:b]] r:−F d[e[r:a]]",
    r"\oint\vec{B}\cdot d\vec{l}=\mu_{0}I":
        "nary{chr=∮,limLoc=subSup,subHide=1,supHide=1}[sub[] sup[] e[acc{chr=\u20d7}[e[r:B]] r:⋅d "
        "acc{chr=\u20d7}[e[r:l]]]] r:= sSub[e[r:μ] sub[r:0]] r:I",
    # delimiters: \left..\right, paired bare brackets, half-open intervals, |..|
    r"\left.F(x)\right|_{a}^{b}": "sSubSup[e[d{begChr=,endChr=|}[e[r:F d[e[r:x]]]]] sub[r:a] sup[r:b]]",
    r"x\in\left[1;3\right)": "r:x∈ d{begChr=[}[e[r:1;3]]",
    r"x\in(0;+\infty)": "r:x∈ d[e[r:0;+∞]]",
    r"D=\mathbb{R}\setminus\{1\}": "r:D=ℝ∖ d{begChr={,endChr=}}[e[r:1]]",
    r"I=\frac{\mathcal{E}}{R_{N}+r}": "r:I= f[num[r:ℰ] den[sSub[e[r:R] sub[r:N]] r:+r]]",
    r"\mathbb{N}^{*}+\mathfrak{g}+\mathsf{A1}": "sSup[e[r:ℕ] sup[r:*]] r:+𝔤+𝖠𝟣",
    r"F=k\frac{|q_{1}q_{2}|}{\varepsilon r^{2}}":
        "r:F=k f[num[d{begChr=|,endChr=|}[e[sSub[e[r:q] sub[r:1]] sSub[e[r:q] sub[r:2]]]]] "
        "den[r:ε sSup[e[r:r] sup[r:2]]]]",
    r"\vec{a}\cdot\vec{b}=|\vec{a}|.|\vec{b}|\cos(\vec{a},\vec{b})":
        "acc{chr=\u20d7}[e[r:a]] r:⋅ acc{chr=\u20d7}[e[r:b]] r:= d{begChr=|,endChr=|}[e[acc{chr=\u20d7}[e[r:a]]]] "
        "r:. d{begChr=|,endChr=|}[e[acc{chr=\u20d7}[e[r:b]]]] "
        "func[fName[r@p:cos] e[d[e[acc{chr=\u20d7}[e[r:a]] r:, acc{chr=\u20d7}[e[r:b]]]]]]",
    r"\left(\frac{a}{b}\middle|c\right)": "d[e[f[num[r:a] den[r:b]]] e[r:c]]",
    r"\left\lfloor x\right\rfloor+\lceil y\rceil": "d{begChr=⌊,endChr=⌋}[e[r:x]] r:+ d{begChr=⌈,endChr=⌉}[e[r:y]]",
    # prescripts
    r"{}_{92}^{235}\mathrm{U}+{}_{0}^{1}n\to{}_{56}^{141}\mathrm{Ba}+{}_{36}^{92}\mathrm{Kr}+3{}_{0}^{1}n":
        "sPre[sub[r:92] sup[r:235] e[r@p:U]] r:+ sPre[sub[r:0] sup[r:1] e[r:n]] r:→ "
        "sPre[sub[r:56] sup[r:141] e[r@p:Ba]] r:+ sPre[sub[r:36] sup[r:92] e[r@p:Kr]] r:+3 "
        "sPre[sub[r:0] sup[r:1] e[r:n]]",
    # environments
    r"f(x)=\begin{cases}x^{2} & \text{khi } x\ge 0\\-x & \text{khi } x<0\end{cases}":
        "r:f d[e[r:x]] r:= d{begChr={,endChr=}[e[m[mr[e[sSup[e[r:x] sup[r:2]]] e[r~:khi  r:x≥0]] "
        "mr[e[r:−x] e[r~:khi  r:x<0]]]]]",
    r"\left\{\begin{array}{l}x=1+2t\\y=3-t\end{array}\right.":
        "d{begChr={,endChr=}[e[m[mr[e[r:x=1+2t]] mr[e[r:y=3−t]]]]]",
    r"\begin{aligned}x&=1\\y&=2\end{aligned}": "eqArr[e[r:x&=1] e[r:y&=2]]",
    r"A=\begin{pmatrix}1&2\\3&4\end{pmatrix}": "r:A= d[e[m[mr[e[r:1] e[r:2]] mr[e[r:3] e[r:4]]]]]",
    r"\binom{n}{k}=\dbinom{n}{n-k}":
        "d[e[f{type=noBar}[num[r:n] den[r:k]]]] r:= d[e[f{type=noBar}[num[r:n] den[r:n−k]]]]",
    r"\overset{t^{\circ}}{\rightarrow}": "limUpp[e[r:→] lim[r:t°]]",
    r"{a \over b}+1": "f[num[r:a] den[r:b]] r:+1",
}

# Must raise ValueError: malformed or outside the supported subset.
BAD = [
    "", "   ", "$$", "{}", r"\,\quad", r"\frac{1}", r"\frac{1}{2", r"\sqrt", "x^", "x_", "x^2^3", "x_1_2",
    "x''^2^3", r"\foo{x}", "}", "{x", r"\left( x", r"x \right)", r"\left\alpha x\right)", "a & b",
    r"a \\ b", r"\begin{tabular}{c}x\end{tabular}", r"\begin{cases}x\end{matrix}", r"\begin{cases}x",
    r"\begin{array}{p{2cm}}x\end{array}", "x # y", "$x$ $y$", r"\text{$x$}", r"\text x",
    r"\text{a^2}", r"\color{red}{x}", r"\textcolor{red}{x}", "x\x00y", "x\u200by", r"\not\alpha",
    r"\lim_{x}^{y} x", "{" * 60 + "x" + "}" * 60, r"\sqrt" * 60 + "x", "x" * 2001, "x \\", r"\over",
    r"\boxed{x}", r"\underbrace{x}_{n}", r"\operatorname{}", r"\operatorname{\frac{a}{b}}",
    r"\begin{gathered}a&b\end{gathered}", r"a \over b \over c", r"\overbrace{x}", r"\mathrm",
    None, 42, b"x",
]

SHOWN = {"chr", "limLoc", "subHide", "supHide", "begChr", "sepChr", "endChr", "pos", "type"}


def skeleton(el):
    """Compact structure of an OMML element: runs as r:text, structures as tag[children]."""
    parts = []
    for c in el:
        if not c.tag.startswith(M):
            continue
        tag = c.tag[len(M):]
        if tag == "r":
            mark = ""
            rpr = c.find(M + "rPr")
            if rpr is not None:
                sty = rpr.find(M + "sty")
                mark = "~" if rpr.find(M + "nor") is not None else "@" + sty.get(M + "val") if sty is not None else ""
            parts.append(f"r{mark}:" + "".join(t.text or "" for t in c.iter(M + "t")))
        elif not tag.endswith("Pr"):
            pr = c.find(M + tag + "Pr")
            props = [f"{p.tag[len(M):]}={p.get(M + 'val')}" for p in (pr if pr is not None else [])
                     if p.tag[len(M):] in SHOWN]
            parts.append(tag + ("{" + ",".join(props) + "}" if props else "") + f"[{skeleton(c)}]")
    return " ".join(parts)


# Content models of the OMML elements app.omml writes (ECMA-376 shared-math, transitional), as
# regular expressions over the child names; WordprocessingML children are prefixed "w:".
EL = "(?:r|f|rad|sSub|sSup|sSubSup|sPre|d|acc|bar|func|limLow|limUpp|nary|eqArr|m) "
ARG = f"(?:argPr )?(?:{EL})*(?:ctrlPr )?"
CONTENT = {
    "oMath": f"(?:{EL})+",
    "r": "(?:rPr )?(?:w:rPr )?t ",
    "rPr": "(?:lit )?(?:nor |(?:scr )?(?:sty )?)(?:brk )?(?:aln )?",
    "f": "(?:fPr )?num den ", "fPr": "(?:type )?(?:ctrlPr )?",
    "rad": "(?:radPr )?deg e ", "radPr": "(?:degHide )?(?:ctrlPr )?",
    "sSub": "(?:sSubPr )?e sub ", "sSubPr": "(?:ctrlPr )?",
    "sSup": "(?:sSupPr )?e sup ", "sSupPr": "(?:ctrlPr )?",
    "sSubSup": "(?:sSubSupPr )?e sub sup ", "sSubSupPr": "(?:alnScr )?(?:ctrlPr )?",
    "sPre": "(?:sPrePr )?sub sup e ", "sPrePr": "(?:ctrlPr )?",
    "d": "(?:dPr )?(?:e )+", "dPr": "(?:begChr )?(?:sepChr )?(?:endChr )?(?:grow )?(?:shp )?(?:ctrlPr )?",
    "acc": "(?:accPr )?e ", "accPr": "(?:chr )?(?:ctrlPr )?",
    "bar": "(?:barPr )?e ", "barPr": "(?:pos )?(?:ctrlPr )?",
    "func": "(?:funcPr )?fName e ", "funcPr": "(?:ctrlPr )?",
    "limLow": "(?:limLowPr )?e lim ", "limLowPr": "(?:ctrlPr )?",
    "limUpp": "(?:limUppPr )?e lim ", "limUppPr": "(?:ctrlPr )?",
    "nary": "(?:naryPr )?sub sup e ",
    "naryPr": "(?:chr )?(?:limLoc )?(?:grow )?(?:subHide )?(?:supHide )?(?:ctrlPr )?",
    "eqArr": "(?:eqArrPr )?(?:e )+",
    "eqArrPr": "(?:baseJc )?(?:maxDist )?(?:objDist )?(?:rSpRule )?(?:rSp )?(?:ctrlPr )?",
    "m": "(?:mPr )?(?:mr )+",
    "mPr": "(?:baseJc )?(?:plcHide )?(?:rSpRule )?(?:cGpRule )?(?:rSp )?(?:cSp )?(?:cGp )?(?:mcs )?(?:ctrlPr )?",
    "mcs": "(?:mc )+", "mc": "(?:mcPr )?", "mcPr": "(?:count )?(?:mcJc )?", "mr": "(?:e )+",
    "argPr": "(?:scrLvl )?", "ctrlPr": "(?:w:rPr )?",
    "w:rPr": "(?:w:rFonts )?(?:w:b )?(?:w:i )?(?:w:color )?(?:w:sz )?(?:w:szCs )?",
}
for _arg in ("num", "den", "deg", "e", "sub", "sup", "fName", "lim"):
    CONTENT[_arg] = ARG
ONOFF = {"on", "off", "1", "0", "true", "false"}
VALUES = {
    "type": {"bar", "skw", "lin", "noBar"}, "limLoc": {"undOvr", "subSup"}, "pos": {"top", "bot"},
    "sty": {"p", "b", "i", "bi"}, "mcJc": {"left", "center", "right", "inside", "outside"},
    "scr": {"roman", "script", "fraktur", "double-struck", "sans-serif", "monospace"},
    "scrLvl": {"-2", "-1", "0", "1", "2"}, "degHide": ONOFF, "subHide": ONOFF, "supHide": ONOFF,
    "plcHide": ONOFF, "nor": ONOFF,
}
CHAR_VALUES = {"chr", "begChr", "sepChr", "endChr"}  # ST_Char: at most one character
LEAVES = set(VALUES) | CHAR_VALUES | {"count", "t", "w:rFonts", "w:b", "w:i", "w:color", "w:sz", "w:szCs"}


def _name(el):
    return "w:" + el.tag[len(W):] if el.tag.startswith(W) else el.tag[len(M):]


def schema_errors(root):
    errors = []
    for el in root.iter():
        if not (el.tag.startswith(M) or el.tag.startswith(W)):
            errors.append(f"foreign element {el.tag}")
            continue
        name = _name(el)
        kids = "".join(_name(c) + " " for c in el)
        val = el.get(f"{{{M_NS}}}val")
        if name in CONTENT:
            if not re.fullmatch(CONTENT[name], kids):
                errors.append(f"<{name}> children: {kids!r}")
        elif name in LEAVES:
            if len(el):
                errors.append(f"<{name}> must be empty")
            if name in VALUES and val is not None and val not in VALUES[name]:
                errors.append(f"<{name} val={val!r}>")
            if name in CHAR_VALUES and (val is None or len(val) > 1):
                errors.append(f"<{name} val={val!r}> is not one character")
            if name == "count" and not (val or "").isdigit():
                errors.append(f"<count val={val!r}>")
            if name == "t" and not el.text:
                errors.append("empty <t>")
            if name == "w:sz" and not (el.get(f"{{{W_NS}}}val") or "").isdigit():
                errors.append("bad <w:sz>")
        else:
            errors.append(f"unexpected <{name}>")
    for rad in root.iter(M + "rad"):
        hidden = rad.find(f"{M}radPr/{M}degHide") is not None
        if hidden == bool(len(rad.find(M + "deg"))):
            errors.append("degHide does not match the degree")
    for m in root.iter(M + "m"):
        cols = sum(int(c.get(M + "val")) for c in m.findall(f"{M}mPr/{M}mcs/{M}mc/{M}mcPr/{M}count"))
        if any(len(mr) != cols for mr in m.findall(M + "mr")):
            errors.append("matrix row width differs from its column count")
    return errors


class LatexToOmmlTest(unittest.TestCase):
    def test_corpus_converts_to_valid_omml(self):
        corpus = GOLD + MODEL + WRITTEN
        self.assertGreaterEqual(len(WRITTEN), 40)
        for tex in corpus:
            with self.subTest(tex=tex):
                el = latex_to_omml(tex)
                self.assertEqual(el.tag, M + "oMath")
                self.assertEqual(schema_errors(el), [])

    def test_skeletons(self):
        for tex, want in SKELETONS.items():
            with self.subTest(tex=tex):
                self.assertEqual(skeleton(latex_to_omml(tex)), want)

    def test_gold_text_round_trip(self):
        for tex in GOLD:
            with self.subTest(tex=tex):
                self.assertEqual(to_tex(latex_to_omml(tex)), re.sub(r"_\{(.)\}", r"_\1", tex).replace("-", "−"))

    def test_bad_input_raises_value_error(self):
        for tex in BAD:
            with self.subTest(tex=tex):
                with self.assertRaises(ValueError):
                    latex_to_omml(tex)

    def test_random_input_is_valid_or_value_error(self):
        import random

        tokens = ["x", "2", "-", "=", "(", ")", "[", "]", "|", "{", "}", "^", "_", "'", "&", r"\\", "~", ".",
                  r"\frac", r"\sqrt", r"\left(", r"\right)", r"\left.", r"\right|", r"\middle|", r"\sin",
                  r"\lim", r"\sum", r"\int", r"\limits", r"\mathrm", r"\text{a b}", r"\vec", r"\circ", r"\,",
                  r"\!", "{}", r"\begin{cases}", r"\end{cases}", r"\begin{array}{lc}", r"\end{array}", r"\{",
                  r"\}", r"\lvert", r"\rvert", r"\big(", r"\over", r"\binom", r"\mathbb", r"\not", r"\overset"]
        rnd = random.Random(7)
        for _ in range(3000):
            tex = "".join(rnd.choice(tokens) for _ in range(rnd.randint(1, 14)))
            spaces = rnd.choice(["word", "none", "all"])
            try:
                el = latex_to_omml(tex, spaces=spaces)
            except ValueError:
                continue
            self.assertEqual(schema_errors(el), [], tex)

    def test_math_delimiters_and_nfc(self):
        a = skeleton(latex_to_omml(r"\frac{a}{b}"))
        for wrapped in (r"$\frac{a}{b}$", r"$$\frac{a}{b}$$", r"\(\frac{a}{b}\)", r"\[\frac{a}{b}\]"):
            self.assertEqual(skeleton(latex_to_omml(wrapped)), a)
        self.assertEqual(skeleton(latex_to_omml("\\text{ho\u0300}")), "r~:hò")

    def test_run_properties(self):
        el = latex_to_omml(r"x+\text{ khi } y", size_pt=13, text_font="Times New Roman")
        runs = el.findall(M + "r")
        fonts = [r.find(f"{W}rPr/{W}rFonts").get(W + "ascii") for r in runs]
        self.assertEqual(fonts, ["Cambria Math", "Times New Roman", "Cambria Math"])
        self.assertEqual({s.get(W + "val") for s in el.iter(W + "sz")}, {"26"})
        self.assertEqual(runs[1].find(M + "t").get("{http://www.w3.org/XML/1998/namespace}space"), "preserve")
        bold = latex_to_omml(r"\textbf{A}").find(f"{M}r/{W}rPr/{W}b")
        self.assertIsNotNone(bold)

    def test_color(self):
        # a rebuilt equation of a native PDF keeps its colour: w:color in every run, schema order kept
        root = latex_to_omml(r"\frac{a}{b}+\text{là}", size_pt=11, text_font="Cambria Math", color="0000FF")
        self.assertEqual(schema_errors(root), [])
        colors = {c.get(f"{W}val") for c in root.iter(f"{W}color")}
        self.assertEqual(colors, {"0000FF"})
        self.assertEqual(len(list(root.iter(f"{W}color"))), len(list(root.iter(f"{W}rPr"))))
        for bad in ("blue", "00F", "0000FG"):
            with self.assertRaises(ValueError):
                latex_to_omml("x", color=bad)

    def test_spacing(self):
        tex = r"~5\,\mathrm{kg}\quad x\!y~\,z~."
        self.assertEqual(skeleton(latex_to_omml(tex)), "r:5\u00a0 r@p:kg r:\u00a0xy\u00a0z.")
        self.assertEqual(skeleton(latex_to_omml(tex, spaces="none")), "r:5 r@p:kg r:xyz.")
        exact = latex_to_omml(tex, spaces="all")
        self.assertEqual(skeleton(exact), "r:5\u2009 r@p:kg r:\u2003xy\u00a0\u2009z\u00a0.")
        self.assertEqual(schema_errors(exact), [])
        self.assertEqual(skeleton(latex_to_omml(r"\int x\,dx")), "nary{chr=∫,limLoc=subSup,subHide=1,supHide=1}"
                                                               "[sub[] sup[] e[r:x\u00a0dx]]")
        self.assertEqual(skeleton(latex_to_omml(r"\sin\,x")), "func[fName[r@p:sin] e[r:x]]")
        with self.assertRaises(ValueError):
            latex_to_omml("x", spaces="some")

    def test_in_a_docx(self):
        import io

        from docx import Document

        doc = Document()
        para = doc.add_paragraph("A. ")
        para._p.append(latex_to_omml(r"v_{cqp}=\sqrt{\frac{3p}{\rho}}"))
        buf = io.BytesIO()
        doc.save(buf)
        again = Document(io.BytesIO(buf.getvalue()))
        maths = again.paragraphs[0]._p.findall(M + "oMath")
        self.assertEqual(len(maths), 1)
        self.assertEqual(schema_errors(maths[0]), [])


def to_tex(el):
    """LaTeX back from the OMML structures the gold formulas use (fractions, roots, scripts)."""
    out = []
    for c in el:
        tag = c.tag[len(M):] if c.tag.startswith(M) else ""
        get = lambda name: to_tex(c.find(M + name))  # noqa: E731
        if tag == "r":
            out.append("".join(t.text for t in c.iter(M + "t")))
        elif tag == "f":
            out.append(r"\frac{" + get("num") + "}{" + get("den") + "}")
        elif tag == "rad":
            deg = get("deg")
            out.append((r"\sqrt[" + deg + "]{" if deg else r"\sqrt{") + get("e") + "}")
        elif tag == "sSub":
            sub = get("sub")
            out.append(get("e") + ("_" + sub if len(sub) == 1 else "_{" + sub + "}"))
        elif tag == "sSup":
            out.append(get("e") + "^{" + get("sup") + "}")
    return "".join(out).replace("ρ", r"\rho")


# Optional checks against Microsoft Word's own transforms (local only, never shipped).
_WORD = Path("/Applications/Microsoft Word.app/Contents/Resources")
try:
    from latex2mathml.converter import convert as _to_mathml
except ImportError:
    _to_mathml = None
_HAVE_WORD = _to_mathml is not None and (_WORD / "mathml2omml.xsl").is_file()


def _merged(el):
    """Structure with properties dropped and adjacent runs merged (style ignored)."""
    parts = []
    for c in el:
        tag = c.tag[len(M):] if c.tag.startswith(M) else ""
        if not tag or tag.endswith("Pr"):
            continue
        if tag == "r":
            text = "".join(t.text or "" for t in c.iter(M + "t"))
            if parts and parts[-1].startswith("r:"):
                parts[-1] += text
            else:
                parts.append("r:" + text)
        else:
            parts.append(f"{tag}[{_merged(c)}]")
    return " ".join(parts)


@unittest.skipUnless(_HAVE_WORD, "needs Microsoft Word's mathml2omml.xsl and latex2mathml (local only)")
class MicrosoftComparisonTest(unittest.TestCase):
    """Same structure as Word's own MathML -> OMML transform where both should agree."""

    AGREE = GOLD + [
        r"\frac{p_{1}V_{1}}{T_{1}}=\frac{p_{2}V_{2}}{T_{2}}", r"v_{cqp}=\sqrt{\frac{3RT}{\mu}}",
        r"\omega=\frac{1}{\sqrt{LC}}",
        r"\Delta t=\frac{\Delta t_{0}}{\sqrt{1-\frac{v^{2}}{c^{2}}}}", r"x=\frac{-b\pm\sqrt{b^{2}-4ac}}{2a}",
        r"\sqrt[3]{x^{2}}=x^{\frac{2}{3}}", r"v^{2}-v_{0}^{2}=2as", r"N=N_{0}e^{-\lambda t}",
        r"\sqrt{x^{2}+y^{2}}", r"\prod_{i=1}^{n}a_{i}",
        r"\sum_{k=1}^{n}k^{2}", r"\left(\frac{a}{b}\right)^{2}", r"e^{i\pi}+1=0",
    ]

    def test_same_structure(self):
        xslt = etree.XSLT(etree.parse(str(_WORD / "mathml2omml.xsl")))
        for tex in self.AGREE:
            with self.subTest(tex=tex):
                theirs = xslt(etree.fromstring(_to_mathml(tex))).getroot()
                self.assertEqual(_merged(latex_to_omml(tex)), _merged(theirs))


def _mathml(el):
    """Compact MathML tree: tag(children) or tag:text, single-child mrow collapsed."""
    tag, kids = el.tag.split("}")[-1], list(el)
    if not kids:
        return f"{tag}:{(el.text or '').strip()}"
    if tag == "mrow" and len(kids) == 1:
        return _mathml(kids[0])
    return f"{tag}(" + " ".join(_mathml(k) for k in kids) + ")"


@unittest.skipUnless((_WORD / "omml2mathml.xsl").is_file(), "needs Microsoft Word's omml2mathml.xsl (local only)")
class WordReadsBackTest(unittest.TestCase):
    """Word's own OMML -> MathML transform reads every corpus equation as intended."""

    EXPECT = {
        r"v_{cqp}=\sqrt{\frac{3p}{\rho}}": "math(msub(mi:v mrow(mi:c mi:q mi:p)) mo:= msqrt(mfrac(mrow(mn:3 mi:p) mi:ρ)))",
        r"\sum_{k=1}^{n}k": "math(mrow(munderover(mo:∑ mrow(mi:k mo:= mn:1) mi:n) mi:k))",
        r"\sin x": "math(mrow(mi:sin mo:\u2061 mi:x))",
        r"{}_{92}^{235}\mathrm{U}": "math(mmultiscripts(mi:U mprescripts: mn:92 mn:235))",
        r"\vec{F}": "math(mover(mi:F mo:→))",
    }

    def test_corpus(self):
        xslt = etree.XSLT(etree.parse(str(_WORD / "omml2mathml.xsl")))
        for tex in GOLD + MODEL + WRITTEN:
            with self.subTest(tex=tex):
                mml = xslt(etree.fromstring(etree.tostring(latex_to_omml(tex)))).getroot()
                self.assertTrue(len(mml))
                if tex in self.EXPECT:
                    self.assertEqual(_mathml(mml), self.EXPECT[tex])
        for tex, want in self.EXPECT.items():
            self.assertEqual(_mathml(xslt(etree.fromstring(etree.tostring(latex_to_omml(tex)))).getroot()), want)


if __name__ == "__main__":
    unittest.main()
