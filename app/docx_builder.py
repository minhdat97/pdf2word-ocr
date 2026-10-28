"""Write page layouts (native text or OCR) to an editable .docx."""

import io
import re
from dataclasses import replace
from itertools import count, pairwise
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.exceptions import InvalidSpanError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .model import LINE_BREAK, TAB, Columns, PageLayout, Paragraph, Picture, Run, Table
from .omml import latex_to_omml

FONT = "Times New Roman"
EMU_PER_PT = 12700
BOTTOM_SLACK = 10.0  # pt below the lowest line of the fullest page
TAB_ALIGN = {
    "left": WD_TAB_ALIGNMENT.LEFT,
    "center": WD_TAB_ALIGNMENT.CENTER,
    "right": WD_TAB_ALIGNMENT.RIGHT,
}
# Word's highlight colours (w:highlight); other fills are written as shading (w:shd).
HIGHLIGHTS = {
    "FFFF00": "yellow", "00FF00": "green", "00FFFF": "cyan", "FF00FF": "magenta", "0000FF": "blue",
    "FF0000": "red", "000080": "darkBlue", "008080": "darkCyan", "008000": "darkGreen",
    "800080": "darkMagenta", "800000": "darkRed", "808000": "darkYellow", "808080": "darkGray",
    "C0C0C0": "lightGray",
}
ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _clamp(value: float, lo: float, hi: float) -> float:
    return max(lo, min(hi, value))


def _margins(pages: list[PageLayout], header: list[Paragraph], footer: list[Paragraph]):
    filled = [p for p in pages if p.elements]
    if not filled:
        return 72.0, 72.0, 72.0, 72.0
    left = min(e.x0 for p in filled for e in p.elements)
    # Body lines end where at least two blocks reach: a lone line the PDF let run
    # into the margin (a justified line with a formula, 567 against 556pt) would
    # otherwise widen every line, and Word would fit more on each and on each page.
    edges = sorted((e.x1 for p in filled for e in p.elements), reverse=True)
    edge = next((x for x in edges if sum(y >= x - 1.5 for y in edges) >= 2), edges[0])
    right = pages[0].width_pt - edge
    top = min(p.elements[0].y0 for p in filled)
    # Room for Word setting lines a point or two taller than the PDF; where pages
    # flow on, the paragraphs opening a page are kept together against it.
    bottom = min(p.height_pt - max(e.y1 for e in p.elements) for p in filled) - BOTTOM_SLACK
    if header:
        top = max(top, max(h.y1 for h in header) + 2)
    if footer:
        bottom = max(bottom, pages[0].height_pt - min(f.y0 for f in footer) + 2)
    # Top/bottom may be tiny: content placed near the paper edge must still fit.
    return (_clamp(left, 14, 108), _clamp(right, 14, 108), _clamp(top, 4, 108), _clamp(bottom, 4, 108))


# --- runs and paragraphs -------------------------------------------------------


def _format(run, r: Run) -> None:
    if r.size:
        run.font.size = Pt(r.size)
    run.bold = r.bold or None
    run.italic = r.italic or None
    if r.underline:
        run.underline = True
    if r.color:
        run.font.color.rgb = RGBColor.from_string(r.color)
    font = _xml_text(r.font) if r.font else None  # a PDF's font name may hold any byte
    if font:
        run.font.name = font
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    if r.script == "super":
        run.font.superscript = True
    elif r.script == "sub":
        run.font.subscript = True
    if r.highlight:
        rpr = run._r.get_or_add_rPr()
        if r.highlight in HIGHLIGHTS:
            mark = OxmlElement("w:highlight")
            mark.set(qn("w:val"), HIGHLIGHTS[r.highlight])
        else:
            mark = OxmlElement("w:shd")
            mark.set(qn("w:val"), "clear")
            mark.set(qn("w:color"), "auto")
            mark.set(qn("w:fill"), r.highlight)
        _insert_rpr(rpr, mark)


# Schema order of the rPr children python-docx may already have written.
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
    "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
    "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
    "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
    "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
    "textDirection", "tcFitText", "vAlign", "hideMark",
]


def _insert_ordered(parent, el, order: list[str]) -> None:
    """Insert `el` into `parent` at its schema position (Word rejects out-of-order children)."""
    rank = order.index(el.tag.split("}")[1])
    for child in parent:
        name = child.tag.split("}")[1]
        if name in order and order.index(name) > rank:
            child.addprevious(el)
            return
    parent.append(el)


def _insert_rpr(rpr, el) -> None:
    _insert_ordered(rpr, el, RPR_ORDER)


FORMULA_ALT = "Công thức (ảnh)"


def _add_formula(p, r: Run) -> None:
    """A formula kept as a picture, inline on the text line: lowered by its depth
    below the baseline, with alt text marking it as a formula."""
    run = p.add_run()
    run.add_picture(io.BytesIO(r.image), width=Pt(r.image_w), height=Pt(r.image_h))
    if r.image_depth > 0.2:
        pos = OxmlElement("w:position")
        pos.set(qn("w:val"), str(-round(r.image_depth * 2)))  # half-points, negative lowers
        _insert_rpr(run._r.get_or_add_rPr(), pos)
    doc_pr = run._r.find(qn("w:drawing")).find(qn("wp:inline")).find(qn("wp:docPr"))
    if doc_pr is not None:
        doc_pr.set("descr", FORMULA_ALT)
        doc_pr.set("title", FORMULA_ALT)


# Characters an XML document cannot hold; a PDF's text layer may give them (U+FFFF from a broken map).
NOT_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _xml_text(text: str) -> str:
    return NOT_XML.sub("", text)


def _add_math(p, r: Run, size: float) -> None:
    """A formula as a Word equation (m:oMath in the paragraph, between its runs). LaTeX the
    converter rejects falls back to the run's picture, else to its text."""
    try:
        # a rebuilt Word equation (native PDF) keeps its colour and writes its words in its own font
        math = latex_to_omml(r.math, size_pt=r.size or size, text_font=_xml_text(r.font or FONT), color=r.color)
        if r.display:  # a display equation: its paragraph is the m:oMathPara (aligned by _align_display)
            para = OxmlElement("m:oMathPara")
            jc = OxmlElement("m:jc")
            jc.set(qn("m:val"), "left")
            para.append(OxmlElement("m:oMathParaPr"))
            para[0].append(jc)
            para.append(math)
            math = para
        p._p.append(math)
    except ValueError:
        if r.image:
            _add_formula(p, r)
        else:
            run = p.add_run(_xml_text(r.math_text if r.math_text is not None else r.math))
            _format(run, r)


def _align_display(p, align: str | None) -> None:
    """A display equation follows its paragraph's alignment (Word centres one by default)."""
    for jc in p._p.iter(qn("m:jc")):
        jc.set(qn("m:val"), {"center": "center", "right": "right"}.get(align or "", "left"))


def _add_field(p, r: Run) -> None:
    # The run text is the value shown until Word computes the field (viewers
    # such as QuickLook never do), e.g. the real page count for NUMPAGES.
    run = p.add_run(_xml_text(r.text) or "1")
    _format(run, r)
    fld = OxmlElement("w:fldSimple")
    # MERGEFORMAT keeps the run's formatting when Word recomputes the field.
    fld.set(qn("w:instr"), f"{r.field} \\* MERGEFORMAT")
    p._p.remove(run._r)
    fld.append(run._r)
    p._p.append(fld)


def _write_runs(p, runs: list[Run], text: str, size: float, bold: bool) -> None:
    if not runs:
        run = p.add_run(_xml_text(text))
        run.font.size = Pt(size)
        run.bold = bold or None
        return
    for r in runs:
        if r.field:
            _add_field(p, r)
        elif r.math:
            _add_math(p, r, size)
        elif r.image:
            _add_formula(p, r)
        elif r.text == TAB:
            p.add_run().add_tab()
        else:
            parts = r.text.split(LINE_BREAK)
            for i, part in enumerate(parts):
                run = p.add_run(_xml_text(part))
                _format(run, r)
                if i < len(parts) - 1:
                    run.add_break()
    if any(r.math and r.font and not r.display for r in runs):  # rebuilt equations of native pages (OCR ones have no font)
        _keep_inline(p, max((r.size or size for r in runs if r.math), default=size))


def _keep_inline(p, size: float) -> None:
    """Word sets an equation that fills a paragraph, or a line of it between line breaks, as a display
    equation: centred, with full-size fractions. A space after it keeps it in the line, as the PDF has it."""
    segment: list = []
    for child in [*p._p, None]:
        breaks = child is not None and child.tag == qn("w:r") and (
            child.find(qn("w:br")) is not None or child.find(qn("w:cr")) is not None)
        if child is not None and not breaks:
            segment.append(child)
            continue
        maths = [el for el in segment if el.tag == qn("m:oMath")]
        filled = any(
            (el.tag == qn("w:r") and ((el.find(qn("w:t")) is not None and el.find(qn("w:t")).text)
                                      or el.find(qn("w:tab")) is not None or el.find(qn("w:drawing")) is not None))
            or el.tag == qn("w:fldSimple") for el in segment)
        if maths and not filled:
            space = p.add_run(" ")
            space.font.size = Pt(size)
            maths[-1].addnext(space._r)
        segment = []


def _fill_paragraph(
    p, el: Paragraph, origin_x: float, right_edge: float, space_before: float, page_break: bool = False
):
    _write_runs(p, el.runs, el.text, el.font_pt, el.bold)
    fmt = p.paragraph_format
    fmt.alignment = ALIGN[el.align]
    _align_display(p, el.align)
    fmt.space_before = Pt(round(max(space_before, 0), 1))
    if page_break:
        fmt.page_break_before = True
    if el.align in ("left", "justify"):
        # x0 is the leftmost line; undo a hanging indent to get the body edge.
        body_left = el.x0 - min(el.first_line_indent_pt, 0)
        if body_left - origin_x > 1.5:
            fmt.left_indent = Pt(round(body_left - origin_x, 1))
        if el.first_line_indent_pt:
            fmt.first_line_indent = Pt(round(el.first_line_indent_pt, 1))
    if el.natural_right and el.natural_right > right_edge - 0.3:
        # Producers that squeeze spaces fit more on a line than Word will;
        # let the paragraph reach into the margin so line breaks stay put.
        fmt.right_indent = Pt(-round(el.natural_right - right_edge + 0.8, 1))
    for x, kind in el.tab_stops:
        # Tab stops are measured from the margin (or cell edge), not the indent.
        fmt.tab_stops.add_tab_stop(Pt(round(x - origin_x, 1)), TAB_ALIGN[kind])
    if el.line_height_pt:
        fmt.line_spacing = Pt(round(el.line_height_pt, 1))
        # "At least" keeps the PDF pitch without clipping; pitches tighter than
        # Word's single spacing (~1.15em) need "exactly" to be reproduced.
        tight = el.line_height_pt < 1.12 * el.font_pt
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY if tight else WD_LINE_SPACING.AT_LEAST
    elif el.line_pitch_pt and el.line_pitch_pt > 1.4 * el.font_pt:
        # Word's single spacing is ~1.15em for Times; only keep clearly wider spacing.
        fmt.line_spacing = round(el.line_pitch_pt / (1.15 * el.font_pt), 2)
    return p


# --- tables ----------------------------------------------------------------------


def _twips(pt: float) -> str:
    return str(int(round(pt * 20)))


def _set_table_props(
    table, width_pt: float, indent_pt: float, borders: bool, cell_margin: float | None = None
):
    """Direct width, fixed layout, indent and borders: viewers that ignore
    table styles (QuickLook, some web previews) still draw the grid."""
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout", "w:tblBorders", "w:tblCellMar", "w:tblInd"):
        for old in tbl_pr.findall(qn(tag)):
            tbl_pr.remove(old)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), _twips(width_pt))
    width.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    edges = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single" if borders else "nil")
        if borders:
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "000000")
        edges.append(b)
    for el in (width, layout, edges):
        _insert_ordered(tbl_pr, el, TBLPR_ORDER)
    if cell_margin is not None:
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), _twips(cell_margin))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        _insert_ordered(tbl_pr, mar, TBLPR_ORDER)
    if indent_pt > 1.5:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), _twips(indent_pt))
        ind.set(qn("w:type"), "dxa")
        _insert_ordered(tbl_pr, ind, TBLPR_ORDER)


def _add_table(doc, el: Table, margin_left: float):
    table = doc.add_table(rows=len(el.cells), cols=len(el.cells[0]))
    table.style = "Table Grid"
    table.autofit = False
    _set_table_props(table, sum(el.col_widths_pt), el.x0 - margin_left, borders=el.borders, cell_margin=el.cell_margin_pt)
    for c, width in enumerate(el.col_widths_pt):
        table.columns[c].width = Pt(width)
    for r, row in enumerate(el.cells):
        table.rows[r].height = Pt(el.row_heights_pt[r])
        table.rows[r].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for c in range(len(row)):
            table.cell(r, c).width = Pt(el.col_widths_pt[c])
    # Merged cells (widths add up); their text sits in the first grid cell.
    covered: set[tuple[int, int]] = set()
    for r0, c0, rs, cs in el.merges:
        area = {(r, c) for r in range(r0, r0 + rs) for c in range(c0, c0 + cs)}
        if area & covered:
            continue
        try:
            table.cell(r0, c0).merge(table.cell(r0 + rs - 1, c0 + cs - 1))
        except InvalidSpanError:
            continue
        covered |= area - {(r0, c0)}
    for r, row in enumerate(el.cells):
        for c, cell in enumerate(row):
            if (r, c) in covered:
                continue
            target = table.cell(r, c)
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if cell.shading:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), cell.shading)
                _insert_ordered(target._tc.get_or_add_tcPr(), shd, TCPR_ORDER)
            if cell.align:
                target.paragraphs[0].alignment = ALIGN[cell.align]
            if cell.runs or cell.text:
                _write_runs(target.paragraphs[0], cell.runs, cell.text, el.font_pt, cell.bold)
                _align_display(target.paragraphs[0], cell.align)
    return table


def _add_columns(doc, el: Columns, margin_left: float):
    widths = [b - a for a, b in pairwise(el.bounds)]
    table = doc.add_table(rows=1, cols=len(widths))
    table.autofit = False
    _set_table_props(table, sum(widths), el.bounds[0] - margin_left, borders=False, cell_margin=0)
    for c, (paras, (lo, hi)) in enumerate(zip(el.cells, pairwise(el.bounds), strict=True)):
        table.columns[c].width = Pt(hi - lo)
        cell = table.cell(0, c)
        cell.width = Pt(hi - lo)
        for i, para in enumerate(paras):
            target = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            _fill_paragraph(target, para, lo, hi, para.space_before_pt)
    return table


# --- pictures --------------------------------------------------------------------


def _float_picture(paragraph, pic: Picture, z: int, para_top: float | None = None) -> None:
    """Anchor `pic` to `paragraph`: across from the page edge, and down from the
    page top, or from the paragraph's top (`para_top`, page coordinates) so it
    travels with the paragraph."""
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(pic.data), width=Pt(pic.x1 - pic.x0), height=Pt(pic.y1 - pic.y0))
    inline = run._r.find(qn("w:drawing")).find(qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    attrs = {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0",
        "relativeHeight": str(z), "behindDoc": "1" if pic.behind_text else "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }
    for k, v in attrs.items():
        anchor.set(k, v)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    vertical = ("page", pic.y0) if para_top is None else ("paragraph", pic.y0 - para_top)
    for tag, (rel, offset) in (("wp:positionH", ("page", pic.x0)), ("wp:positionV", vertical)):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", rel)
        off = OxmlElement("wp:posOffset")
        off.text = str(int(offset * EMU_PER_PT))
        pos.append(off)
        anchor.append(pos)
    anchor.append(inline.find(qn("wp:extent")))
    effect = OxmlElement("wp:effectExtent")
    for side in "ltrb":
        effect.set(side, "0")
    anchor.append(effect)
    if pic.wrap and not pic.behind_text:
        wrap = OxmlElement("wp:wrapSquare")
        wrap.set("wrapText", "bothSides")
    else:
        wrap = OxmlElement("wp:wrapNone")
    anchor.append(wrap)
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    inline.getparent().replace(inline, anchor)


def _anchor_for(pic: Picture, placed: list[tuple[float, float, object]]) -> tuple[object, float]:
    """The paragraph a floating picture belongs to, with its top: the first one
    beside the picture's top (a long paragraph further down shares more height
    with it but comes after the text that refers to it), else the last one above
    it (a figure below the text that refers to it), else the page's first one.
    A paragraph beside only the lower part of the picture (the options under a
    graph) is not its anchor: hung above it, the picture lands on earlier text."""
    ordered = sorted(placed, key=lambda t: t[0])
    for y0, y1, p in ordered:
        if min(y1, pic.y1) - max(y0, pic.y0) >= min(8.0, 0.5 * (y1 - y0)) and y0 <= pic.y0 + 8:
            return p, y0
    above = [(p, y0) for y0, _, p in ordered if y0 <= pic.y0]
    return above[-1] if above else (ordered[0][2], ordered[0][0])


def _clear_of_lines(pic: Picture, page: PageLayout) -> Picture:
    """A figure reaching up to 4pt into the line box of the text just above or
    below it (lines it is not beside) fitted between them – moved, or scaled
    down by up to 15% about its centre when it cannot move: Word would otherwise
    wrap that line around the figure, pushing a word onto a new line and every
    line after it down."""
    lo, hi = float("-inf"), float("inf")
    for e in page.elements:
        if not isinstance(e, Paragraph) or not (e.x0 < pic.x1 and pic.x0 < e.x1):
            continue
        if e.y0 < pic.y0 < e.y1 <= pic.y0 + 4:
            lo = max(lo, e.y1 + 0.5)
        if pic.y1 - 4 <= e.y0 < pic.y1 < e.y1:
            hi = min(hi, e.y0 - 0.5)
    if lo == float("-inf") and hi == float("inf"):
        return pic
    h, w = pic.y1 - pic.y0, pic.x1 - pic.x0
    top = max(pic.y0, lo)
    if top + h > hi:
        top = max(hi - h, lo)
    if top + h <= hi:
        return replace(pic, y0=top, y1=top + h)
    scale = (hi - top) / h
    if scale < 0.85:
        return pic
    x0 = pic.x0 + w * (1 - scale) / 2
    return replace(pic, x0=x0, y0=top, x1=x0 + w * scale, y1=top + h * scale)


FLOW_ROOM_LINES = 6  # a page ending further than this above the bottom was ended on purpose


def _room_before(prev: PageLayout, page: PageLayout, content_bottom: float, line: float) -> float | None:
    """The room left at the bottom of `prev` when `page` only began because
    prev was full – its content reaches to within a few lines of the bottom
    (what was left is what kept-together lines did not fit in) – so Word may run
    the text on; None when the page break is the author's: prev ends early, a
    page is empty, or `page` opens with a figure above its text. (Lines repeated
    as headers/footers are in the Word header/footer by now; what is left in
    those zones is body text.)"""
    bottoms = [e.y1 for e in prev.elements] + [p.y1 for p in prev.pictures if p.wrap and not p.behind_text]
    if not bottoms or not page.elements:
        return None
    first = page.elements[0]
    if any(p.wrap and not p.behind_text and p.y0 < first.y0 - 2 for p in page.pictures):
        return None
    room = content_bottom - max(bottoms)
    return max(room, 0.0) if room <= FLOW_ROOM_LINES * line else None


def _keep_chain(page: PageLayout, room: float) -> int:
    """How many opening paragraphs of `page` Word must keep together (and with
    each other) so they do not slip into the `room` left on the previous page:
    the first ones whose heights add up to more than it. None is needed when not
    even the first line fits; a table ends the chain."""
    first = page.elements[0]
    if not isinstance(first, Paragraph) or room < (first.line_height_pt or 1.2 * first.font_pt):
        return 0
    total = 0.0
    for j, el in enumerate(page.elements):
        if not isinstance(el, Paragraph):
            return j
        total += (el.space_before_pt if j else 0.0) + el.y1 - el.y0
        if total > room:
            return j + 1
    return len(page.elements)


# --- header / footer -------------------------------------------------------------


def _fill_story(story, paras: list[Paragraph], margin_left: float, right_edge: float, style) -> None:
    prev_bottom = None
    for i, para in enumerate(paras):
        target = story.paragraphs[0] if i == 0 else story.add_paragraph()
        target.style = style
        gap = 0 if prev_bottom is None else para.y0 - prev_bottom
        _fill_paragraph(target, para, margin_left, right_edge, gap)
        prev_bottom = para.y1


# --- document ---------------------------------------------------------------------


def build_docx(
    pages: list[PageLayout],
    out_path: Path,
    header: list[Paragraph] | None = None,
    footer: list[Paragraph] | None = None,
) -> None:
    header, footer = header or [], footer or []
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    # Lines stay where the PDF put them; no widow/orphan reflow across pages.
    normal.paragraph_format.widow_control = False

    section = doc.sections[0]
    section.page_width = Pt(pages[0].width_pt)
    section.page_height = Pt(pages[0].height_pt)
    left, right, top, bottom = _margins(pages, header, footer)
    section.left_margin, section.right_margin = Pt(left), Pt(right)
    section.top_margin, section.bottom_margin = Pt(top), Pt(bottom)
    right_edge = pages[0].width_pt - right
    if header:
        section.header_distance = Pt(max(min(h.y0 for h in header), 4))
        _fill_story(section.header, header, left, right_edge, normal)
    if footer:
        section.footer_distance = Pt(max(pages[0].height_pt - max(f.y1 for f in footer), 4))
        _fill_story(section.footer, footer, left, right_edge, normal)

    z = count(251658240)

    content_bottom = pages[0].height_pt - bottom
    lines = [p.line_height_pt or 1.2 * p.font_pt for pg in pages for p in pg.elements if isinstance(p, Paragraph)]
    body_line = sorted(lines)[len(lines) // 2] if lines else 14.0
    for i, page in enumerate(pages):
        # A page the PDF only started because the previous one was full flows on
        # from it, so edits (or "Chuẩn hóa đề") never leave a near-empty page; its
        # opening paragraphs are kept together just enough not to slip into the
        # room left on the previous page, so unedited pages still break as in the PDF.
        room = _room_before(pages[i - 1], page, content_bottom, body_line) if i else None
        new_page = i > 0 and room is None
        # Word also has the bottom slack (and a point of its own) to fill.
        chain = _keep_chain(page, room + BOTTOM_SLACK + 1) if room is not None else 0
        first_para = None
        placed: list[tuple[float, float, object]] = []  # (y0, y1, paragraph) of this page
        chained: dict[int, tuple[object, float]] = {}  # id(paragraph) in the keep chain -> the chain's first
        figures = [p for p in page.pictures if p.wrap and not p.behind_text]
        prev_p = None  # the previous element's Word paragraph, if it was one
        for j, el in enumerate(page.elements):
            space_before = (el.y0 - top) if j == 0 and (i == 0 or new_page) else el.space_before_pt
            if isinstance(el, Paragraph):
                gap_top = el.y0 - space_before
                met = [f for f in figures if gap_top - 1 < f.y0 < el.y1] if prev_p is not None and space_before > 1 else []
                # A figure spanning the text column (less than 100pt left beside it)
                if any(f.x1 - f.x0 > right_edge - left - 100 for f in met):
                    # Word puts the next line below it by itself, after any space
                    # before; spacing for the figure's height (here or as the
                    # line above's space after) would be added on top of it.
                    space_before = max(el.y0 - max(f.y1 for f in met), 0.0)
                elif met:
                    # Word moves a line meeting a figure's top edge down to that
                    # edge together with its space before; as the previous
                    # paragraph's space after, the gap stays where the PDF has it.
                    prev_p.paragraph_format.space_after = Pt(round(space_before, 1))
                    space_before = 0.0
                p = _fill_paragraph(doc.add_paragraph(), el, left, right_edge, space_before, new_page)
                prev_p = p
                if j < chain:
                    p.paragraph_format.keep_together = el.y1 - el.y0 <= 8 * body_line
                    p.paragraph_format.keep_with_next = j < chain - 1 or None
                    chained[id(p)] = chained[id(placed[0][2])] if chained else (p, el.y0)
                first_para = first_para or p
                placed.append((el.y0, el.y1, p))
            else:
                if new_page or space_before > 2:
                    spacer = doc.add_paragraph()
                    spacer.paragraph_format.page_break_before = new_page or None
                    spacer.paragraph_format.line_spacing = Pt(max(min(space_before, 72), 1))
                    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                (_add_table if isinstance(el, Table) else _add_columns)(doc, el, left)
                prev_p = None
            new_page = False
        if page.pictures and first_para is None:
            first_para = doc.add_paragraph()
            first_para.paragraph_format.page_break_before = new_page or None
            new_page = False
        for pic in page.pictures:
            # A figure is placed below the top of the paragraph beside it (and so
            # belongs to that question), moving with it when text moves to another
            # page. Backgrounds and thin line art (rules, bars: no text wrap) stay
            # page-relative with the page's first paragraph.
            if pic.wrap and not pic.behind_text and placed:
                pic = _clear_of_lines(pic, page)
                anchor, anchor_top = _anchor_for(pic, placed)
                # Word breaks a keep-with-next chain at a paragraph holding a
                # floating picture: such pictures hang from the chain's first one.
                anchor, anchor_top = chained.get(id(anchor), (anchor, anchor_top))
                _float_picture(anchor, pic, next(z), anchor_top)
            else:
                _float_picture(first_para, pic, next(z))
        if new_page:  # empty page
            doc.add_paragraph().paragraph_format.page_break_before = True

    doc.core_properties.author = "pdf2word-ocr"
    doc.save(out_path)
